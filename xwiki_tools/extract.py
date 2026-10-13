"""Bounded local extraction with physical source locators."""

from __future__ import annotations

import csv
import json
import logging
import os
import stat
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

MAX_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 5000
MAX_CELLS = 1_000_000
MAX_CHARACTERS = 20_000_000
MAX_EXPANDED_WORKBOOK_BYTES = 200 * 1024 * 1024
EXTRACTOR_VERSION = "xwiki-source-text-v1"
TEXT_SUFFIXES = {
    ".md",
    ".markdown",
    ".txt",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".rst",
    ".html",
}


@contextmanager
def _quiet_parser(name: str) -> Iterator[None]:
    logger = logging.getLogger(name)
    previous = logger.level
    logger.setLevel(logging.ERROR)
    try:
        yield
    finally:
        logger.setLevel(previous)


@dataclass(frozen=True)
class Segment:
    text: str
    locator: dict[str, Any]
    heading: str | None = None
    context: str = ""


@dataclass
class Extraction:
    source_sha256: str
    format: str
    segments: list[Segment] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    character_count: int = 0
    extractor_version: str = EXTRACTOR_VERSION

    def append(
        self,
        text: str,
        locator: dict[str, Any],
        heading: str | None = None,
        context: str = "",
    ) -> None:
        cleaned_text = text.encode("utf-8", errors="replace").decode("utf-8")
        cleaned_context = context.encode("utf-8", errors="replace").decode("utf-8")
        cleaned_heading = (
            heading.encode("utf-8", errors="replace").decode("utf-8")
            if heading is not None
            else None
        )
        if (cleaned_text, cleaned_context, cleaned_heading) != (text, context, heading):
            warning = "Invalid Unicode code points in extracted text were replaced; review the original source."
            if warning not in self.warnings:
                self.warnings.append(warning)
        text, context, heading = cleaned_text, cleaned_context, cleaned_heading
        self.character_count += len(text) + len(context)
        if self.character_count > MAX_CHARACTERS:
            raise ValueError("Extracted text exceeds the configured character limit.")
        self.segments.append(Segment(text, locator, heading, context))

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def read_stable_bytes(path: Path) -> bytes:
    with path.open("rb") as handle:
        before = os.fstat(handle.fileno())
        if not stat.S_ISREG(before.st_mode) or not 0 < before.st_size <= MAX_BYTES:
            raise ValueError("Choose a non-empty regular file no larger than 50 MiB.")
        content = handle.read(before.st_size + 1)
        after = os.fstat(handle.fileno())
    if len(content) != before.st_size or (before.st_size, before.st_mtime_ns) != (
        after.st_size,
        after.st_mtime_ns,
    ):
        raise ValueError("Source file changed while it was being read.")
    return content


def detect_format(path: Path, content: bytes) -> str:
    suffix = path.suffix.lower().removeprefix(".")
    if b"%PDF-" in content[:1024]:
        return "pdf"
    if suffix in {"csv", "tsv", "xlsx", "xls"}:
        return suffix
    if path.suffix.lower() in TEXT_SUFFIXES:
        return "text"
    return suffix or "binary"


def _value_text(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return json.dumps(value, ensure_ascii=False, default=str)


def _append_row(
    result: Extraction, sheet: str, number: int, values: list[Any], header: str
) -> str:
    if not any(value is not None and value != "" for value in values):
        return header
    try:
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Install the knowledge-tool dependencies to extract spreadsheets."
        ) from exc
    text = " | ".join(
        f"{get_column_letter(index)}{number}={_value_text(value)}"
        for index, value in enumerate(values, 1)
    )
    if not header:
        header = f"First nonempty row (not inferred column names): {text[:600]}"
    result.append(
        text,
        {
            "sheet": sheet,
            "cell_range": f"A{number}:{get_column_letter(len(values))}{number}",
        },
        sheet,
        header,
    )
    return header


def extract_source(path: Path, *, encoding: str = "utf-8") -> Extraction:
    content = read_stable_bytes(path)
    source_format = detect_format(path, content)
    result = Extraction(sha256(content).hexdigest(), source_format)
    if source_format == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install pypdf to extract PDF text.") from exc
        with _quiet_parser("pypdf"):
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are not supported.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError("PDF exceeds the 5,000-page extraction limit.")
            for page_number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    result.warnings.append(
                        f"Page {page_number} has no extractable text; it may be blank or require OCR."
                    )
                result.append(text, {"page": page_number}, f"Page {page_number}")
        result.warnings.append(
            "PDF text order and table layout require source review; images are not analyzed and OCR is not performed."
        )
    elif source_format in {"csv", "tsv"}:
        reader = csv.reader(
            StringIO(content.decode(encoding), newline=""),
            delimiter="\t" if source_format == "tsv" else ",",
            strict=True,
        )
        header, cells = "", 0
        for row_number, values in enumerate(reader, 1):
            cells += len(values)
            if cells > MAX_CELLS:
                raise ValueError(
                    "Spreadsheet exceeds the one-million-cell extraction limit."
                )
            header = _append_row(result, path.name, row_number, values, header)
    elif source_format == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Install openpyxl to extract XLSX files.") from exc
        with ZipFile(BytesIO(content)) as archive:
            if (
                sum(item.file_size for item in archive.infolist())
                > MAX_EXPANDED_WORKBOOK_BYTES
            ):
                raise ValueError(
                    "Expanded workbook exceeds the 200 MiB extraction limit."
                )
        book = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            cells, formula_seen = 0, False
            for sheet in book:
                header = ""
                if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                    raise ValueError(
                        "Worksheet dimensions exceed the extraction limit."
                    )
                for row_number, row in enumerate(sheet.iter_rows(), 1):
                    cells += len(row)
                    if cells > MAX_CELLS:
                        raise ValueError(
                            "Workbook exceeds the one-million-cell extraction limit."
                        )
                    values = []
                    for cell in row:
                        if cell.data_type == "f":
                            formula_seen = True
                            values.append({"formula": cell.value, "evaluated": False})
                        elif cell.data_type == "e":
                            values.append({"excel_error": cell.value})
                        else:
                            values.append(cell.value)
                    header = _append_row(
                        result, sheet.title, row_number, values, header
                    )
            if formula_seen:
                result.warnings.append(
                    "XLSX formulas are preserved but are not evaluated."
                )
            result.warnings.append(
                "Spreadsheet formatting, charts, and merged-cell meaning require source review."
            )
        finally:
            book.close()
    elif source_format == "xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Install xlrd to extract XLS files.") from exc
        book = xlrd.open_workbook(
            file_contents=content, on_demand=True, logfile=StringIO()
        )
        try:
            cells = 0
            for sheet in book.sheets():
                cells += sheet.nrows * sheet.ncols
                if cells > MAX_CELLS:
                    raise ValueError(
                        "Workbook exceeds the one-million-cell extraction limit."
                    )
                header = ""
                for row_index in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(row_index):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(
                                cell.value, book.datemode
                            ).isoformat()
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = {
                                "excel_error": xlrd.error_text_from_code.get(
                                    cell.value, "unknown"
                                )
                            }
                        else:
                            value = cell.value
                        values.append(value)
                    header = _append_row(
                        result, sheet.name, row_index + 1, values, header
                    )
            result.warnings.append(
                "XLS extraction reads stored values and does not recalculate formulas."
            )
        finally:
            book.release_resources()
    elif source_format == "text":
        text = content.decode(encoding)
        result.append(
            text, {"line_start": 1, "line_end": text.count("\n") + 1}, path.name
        )
    else:
        raise ValueError(f"Unsupported extraction format: {source_format}.")
    if not any(segment.text.strip() for segment in result.segments):
        raise ValueError("No readable text or nonempty cells were found.")
    return result
